"""Build the example audit ledger workbook described in ../workbook.md.

Run:  uv run --with openpyxl python build_workbook.py
Writes audit-ledger-example.xlsx and the sample Stage 1 CSVs in csv/.

Formulas use only SUMIFS / COUNTIFS / INDEX / MATCH / IFERROR so the file works in every
Excel version and can be verified outside Excel. FILTER / UNIQUE / VSTACK shorthands from
workbook.md are noted on the README sheet.
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

HERE = Path(__file__).parent
OUT = HERE / "audit-ledger-example.xlsx"
CSV_DIR = HERE / "csv"

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT)
BOLD = Font(name=FONT, bold=True)
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
INPUT_FILL = PatternFill("solid", fgColor="FFFFCC")
MONEY = '#,##0.00;(#,##0.00);"-"'
PCT = "0.0%"

N_LINES = 500     # helper formulas on Lines rows 2..N_LINES
N_DOCS = 50
N_TRACK = 50
N_DETAIL = 100
N_REVIEW = 100
TOL = "Settings!$B$2"

CSV_COLS = ["row_kind", "project", "vendor", "doc_type", "doc_id", "doc_date", "ref_quote_id",
            "page", "line_no", "label_as_printed", "description", "qty", "unit_amount", "amount",
            "category", "confidence", "note", "pages_total", "pages_read", "native_or_scan",
            "mapping_provided", "source_file"]
HELPERS = ["match_key", "labour_group", "invoice_amount", "invoice_hits", "quote_hits", "status",
           "detail_seq", "review_reason", "review_seq", "coverage_seq", "pair_seq"]

# ---------------------------------------------------------------- sample data (generic names)
def doc(project, vendor, doc_type, doc_id, doc_date, ref, pages_total, pages_read, kind, mapping,
        lines, totals, cov_note=""):
    src = f"{project}__{vendor}__{doc_type}__{doc_id}.pdf"
    rows = []
    for i, (page, label, desc, qty, unit, amt, cat, conf, note) in enumerate(lines, 1):
        rows.append(["line", project, vendor, doc_type, doc_id, doc_date, ref, page, i, label, desc,
                     qty, unit, amt, cat, conf, note, "", "", "", "", src])
    for page, label, amt, cat in totals:
        rows.append(["total", project, vendor, doc_type, doc_id, doc_date, ref, page, "", label, "",
                     "", "", amt, cat, "H", "", "", "", "", "", src])
    rows.append(["coverage", project, vendor, doc_type, doc_id, doc_date, ref, "", "", "", "", "", "",
                 "", "", "", cov_note, pages_total, pages_read, kind, mapping, src])
    return doc_id, rows

DOCS = [
    doc("Project Alpha", "Vendor A", "quote", "Q-1001", "2026-05-04", "", 3, "1-3", "native", "Y",
        [(1, "Task chair TC-200", "Task chair, model TC-200", 40, 350, 14000, "PRODUCT", "H", ""),
         (1, "Desk HA-1600", "Height-adjustable desk 1600", 40, 620, 24800, "PRODUCT", "H", ""),
         (2, "Pedestal", "Storage pedestal", 40, 180, 7200, "PRODUCT", "H", ""),
         (2, "Installation", "Installation, 40 workstations, fixed price", "", "", 9600, "INSTALL", "H", ""),
         (2, "Project management", "Project management fee", "", "", 2400, "LABOUR", "H", ""),
         (2, "Freight", "Freight to site", "", "", 1800, "LABOUR", "H", ""),
         (3, "Contingency", "Contingency allowance", "", "", 2000, "ALLOWANCE", "H", ""),
         (3, "Sales tax", "Sales tax 8%", "", "", 3680, "EXCLUDED", "H", "")],
        [(3, "Product subtotal", 46000, "PRODUCT"), (3, "Installation", 9600, "INSTALL"),
         (3, "Services subtotal", 4200, "LABOUR"), (3, "Grand total", 65480, "GRAND")]),
    doc("Project Alpha", "Vendor A", "invoice", "INV-2001", "2026-07-15", "Q-1001", 2, "1-2", "native", "Y",
        [(1, "Task chair TC-200", "Task chair, model TC-200", 40, 350, 14000, "PRODUCT", "H", ""),
         (1, "Desk HA-1600", "Height-adjustable desk 1600", 40, 620, 24800, "PRODUCT", "H", ""),
         (1, "Pedestal", "Storage pedestal", 40, 180, 7200, "PRODUCT", "H", ""),
         (2, "Installation", "Installation, 40 workstations, fixed price", "", "", 9600, "INSTALL", "H", ""),
         (2, "Install labour - additional hours", "Additional installation hours, 32 hrs at 95", 32, 95, 3040, "INSTALL", "M", "matched by principle"),
         (2, "Project management", "Project management fee", "", "", 2400, "LABOUR", "H", ""),
         (2, "Freight", "Freight to site", "", "", 2100, "LABOUR", "H", ""),
         (2, "Site services", "Site services", "", "", 650, "UNMAPPED", "L", "no mapping rule"),
         (2, "Sales tax", "Sales tax 8%", "", "", 3680, "EXCLUDED", "H", "")],
        [(2, "Product subtotal", 46000, "PRODUCT"), (2, "Installation", 12640, "INSTALL"),
         (2, "Services subtotal", 4500, "LABOUR"), (2, "Grand total", 67470, "GRAND")]),
    doc("Project Alpha", "Vendor B", "quote", "Q-1002", "2026-05-10", "", 5, "1-4", "scan", "Y",
        [(1, "Meeting table MT-3200", "Meeting table 3200", 4, 2900, 11600, "PRODUCT", "H", ""),
         (2, "Conference chair", "Conference chair", 32, 410, 13120, "PRODUCT", "M", "scan, figure clear"),
         (2, "Credenza", "Credenza", 4, 1450, 5800, "PRODUCT", "L", "figure partly legible"),
         (3, "Labour", "Delivery and installation", "", "", 4200, "INSTALL", "M", "matched by principle")],
        [(4, "Product subtotal", 30620, "PRODUCT"), (4, "Grand total", 34820, "GRAND")],
        cov_note="page 5 illegible"),
]

MAPPING = [
    ("Installation", "INSTALL", "any", ""),
    ("Install labour", "INSTALL", "any", ""),
    ("Delivery & install", "INSTALL", "any", ""),
    ("Labour", "INSTALL", "any", "Confirm per vendor; some print labour for design time"),
    ("Design fee", "LABOUR", "any", ""),
    ("Project management", "LABOUR", "any", ""),
    ("Freight", "LABOUR", "any", ""),
    ("Allowance", "ALLOWANCE", "any", ""),
    ("Contingency", "ALLOWANCE", "any", ""),
    ("Sales tax", "EXCLUDED", "any", ""),
    ("Surcharge", "EXCLUDED", "any", ""),
]

TRACKER_NOTES = [
    ("Project Alpha", "Vendor A", "LABOUR_ALL", "Invoiced above quoted. See Detail: 2 lines, additional installation hours and a higher freight charge."),
    ("Project Alpha", "Vendor A", "INSTALL", "Invoiced above quoted. See Detail: 1 invoice-only line, installation hours billed by the hour."),
    ("Project Alpha", "Vendor A", "PRODUCT", "matched exactly"),
    ("Project Alpha", "Vendor A", "ALLOWANCE", "Invoiced below quoted. See Detail: 1 quote-only line, contingency not drawn."),
]
DETAIL_NOTES = [
    ("Project Alpha", "Vendor A", "INV-2001", 5, "1", "Hourly installation labour billed on top of the fixed-price installation line."),
    ("Project Alpha", "Vendor A", "Q-1001", 6, "NONE", "Freight invoiced above the quoted amount."),
]
STANDARDS = [
    "Labour is billed at the quoted fixed rate unless a signed change order authorises time and materials.",
    "Installation cost is itemised as a distinct line from product cost.",
    "Product cost is invoiced consistent with the manufacturer-confirmed quote.",
    "Allowance items require documented approval before being invoiced.",
]

# ---------------------------------------------------------------- helpers
def header(ws, cols, row=1):
    for c, name in enumerate(cols, 1):
        cell = ws.cell(row=row, column=c, value=name)
        cell.font = BOLD
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(row=row + 1, column=1)

def widths(ws, w):
    for i, width in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

def put(ws, ref, value, font=BLACK, fmt=None):
    cell = ws[ref]
    cell.value = value
    cell.font = font
    if fmt:
        cell.number_format = fmt
    return cell

def col(name, cols):
    return get_column_letter(cols.index(name) + 1)

# ---------------------------------------------------------------- build
wb = Workbook()
wb.remove(wb.active)

# README -----------------------------------------------------------------
ws = wb.create_sheet("README")
readme = [
    ("Audit ledger — example workbook", BOLD),
    ("Built from workbook.md on 2026-09-13 with generic names. Sample data only: Project Alpha, Vendor A, Vendor B.", BLACK),
    ("", BLACK),
    ("How to read the colours", BOLD),
    ("Blue text on a yellow fill = data you paste or type. Black text = formulas; do not edit.", BLACK),
    ("", BLACK),
    ("Sheets", BOLD),
    ("Lines      Every row from every Stage 1 CSV, pasted under the last row. Columns A–V are the CSV; W onward are helper formulas that must be filled down to the last row.", BLACK),
    ("Mapping    The label table. Paste this whole range into every Stage 1 chat. Add a row when Review shows 'no mapping rule'.", BLACK),
    ("Docs       One row per document, derived from the coverage rows. Stated totals vs extracted totals, sum checks, page coverage.", BLACK),
    ("Tracker    One row per project and vendor. Quoted vs invoiced by category, variance, discrepancy flag.", BLACK),
    ("Detail     Every line that differs between a quote and its invoices: quote only, invoice only, or amount differs.", BLACK),
    ("Review     Everything a human should look at. Empty means no PDF needs opening.", BLACK),
    ("Notes      Where Stage 3 output is pasted. Tracker and Detail look notes up from here.", BLACK),
    ("Settings   The tolerance for 'matches exactly', and the client's billing standards for Stage 3.", BLACK),
    ("", BLACK),
    ("Adding a document", BOLD),
    ("1. Open the CSV the model returned. 2. Copy its data rows (not the header). 3. Paste under the last row of Lines. 4. Fill the helper formulas (W onward) down to the new last row. Everything else updates.", BLACK),
    ("Phase B: Data > Get Data > From Folder on csv/ loads Lines automatically; keep the helper columns beside the query table.", BLACK),
    ("", BLACK),
    ("What the sample data shows", BOLD),
    ("Vendor A: installation billed by the hour on top of the fixed price (invoice-only line), freight above quote (amount differs), contingency not drawn (quote only), one label no rule fits (UNMAPPED, in Review).", BLACK),
    ("Vendor B: quote only, no invoice yet, so Tracker says pending. Scan with one unread page and a stated product total that does not match its lines: both in Review.", BLACK),
    ("", BLACK),
    ("Compatibility note", BOLD),
    ("workbook.md sketches Docs, Tracker, Detail and Review with FILTER, UNIQUE and VSTACK. This file uses the equivalent INDEX/MATCH helper-column pattern so it works in every Excel version and could be verified outside Excel. Same results, longer formulas.", BLACK),
]
for i, (text, font) in enumerate(readme, 1):
    put(ws, f"A{i}", text, font)
ws.column_dimensions["A"].width = 160

# Settings ---------------------------------------------------------------
ws = wb.create_sheet("Settings")
put(ws, "A1", "Setting", BOLD); put(ws, "B1", "Value", BOLD)
put(ws, "A2", "Tolerance (amounts within this are 'matched exactly')", BLACK)
c = put(ws, "B2", 0.01, BLUE, "0.00"); c.fill = INPUT_FILL
put(ws, "A4", "Billing standards (paste into the Stage 3 chat)", BOLD)
for i, s in enumerate(STANDARDS, 1):
    put(ws, f"A{4+i}", f"{i}. {s}", BLUE).fill = INPUT_FILL
widths(ws, [90, 12])
wb.defined_names["Tolerance"] = DefinedName("Tolerance", attr_text="Settings!$B$2")

# Lines ------------------------------------------------------------------
ws = wb.create_sheet("Lines")
cols = CSV_COLS + HELPERS
header(ws, cols)
r = 2
for _, rows in DOCS:
    for row in rows:
        for cidx, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=cidx, value=(v if v != "" else None))
            cell.font = BLUE
            if CSV_COLS[cidx - 1] in ("qty", "unit_amount", "amount") and v != "":
                cell.number_format = MONEY if CSV_COLS[cidx - 1] != "qty" else "0"
        r += 1
last_data_row = r - 1

L = {n: col(n, cols) for n in cols}
A, B, C, D, E, G, K, N, O, P, R_, S = (L[x] for x in ["row_kind", "project", "vendor", "doc_type", "doc_id",
                                                   "ref_quote_id", "description", "amount", "category",
                                                   "confidence", "pages_total", "pages_read"])
W, X, Y, Z, AA, AB, AC, AD, AE, AF, AG = (L[x] for x in HELPERS)
rng = lambda c: f"${c}$2:${c}${N_LINES}"

for i in range(2, N_LINES + 1):
    f = {}
    f[W] = f'=IF({A}{i}="line",IF({D}{i}="quote",{E}{i},{G}{i})&"|"&LOWER(TRIM({K}{i})),"")'
    f[X] = f'=IF(OR({O}{i}="LABOUR",{O}{i}="INSTALL"),"LABOUR_ALL",IF({O}{i}="","",{O}{i}))'
    f[Y] = (f'=IF(AND({A}{i}="line",{D}{i}="quote"),SUMIFS({rng(N)},{rng(W)},{W}{i},{rng(D)},"invoice",'
            f'{rng(A)},"line"),"")')
    f[Z] = (f'=IF(AND({A}{i}="line",{D}{i}="quote"),COUNTIFS({rng(W)},{W}{i},{rng(D)},"invoice",'
            f'{rng(A)},"line"),"")')
    f[AA] = (f'=IF(AND({A}{i}="line",{D}{i}="invoice"),COUNTIFS({rng(W)},{W}{i},{rng(D)},"quote",'
             f'{rng(A)},"line"),"")')
    f[AB] = (f'=IF({A}{i}<>"line","",IF({D}{i}="quote",IF({Z}{i}=0,"quote only",'
             f'IF(ABS({Y}{i}-{N}{i})>{TOL},"amount differs","match")),'
             f'IF({AA}{i}=0,"invoice only","match")))')
    f[AC] = (f'=IF(AND({A}{i}="line",{AB}{i}<>"match",COUNTIFS({rng(A)},"coverage",{rng(B)},{B}{i},'
             f'{rng(C)},{C}{i},{rng(D)},"invoice")>0),MAX(${AC}$1:{AC}{i-1})+1,"")')
    f[AD] = (f'=IF({A}{i}<>"line","",IF({O}{i}="UNMAPPED","no mapping rule",'
             f'IF({P}{i}="L","low confidence","")))')
    f[AE] = f'=IF({AD}{i}<>"",MAX(${AE}$1:{AE}{i-1})+1,"")'
    f[AF] = f'=IF({A}{i}="coverage",MAX(${AF}$1:{AF}{i-1})+1,"")'
    f[AG] = (f'=IF({A}{i}="line",IF(COUNTIFS(${B}$1:{B}{i-1},{B}{i},${C}$1:{C}{i-1},{C}{i},'
             f'${A}$1:{A}{i-1},"line")=0,MAX(${AG}$1:{AG}{i-1})+1,""),"")')
    for cl, formula in f.items():
        cell = ws[f"{cl}{i}"]
        cell.value = formula
        cell.font = BLACK
        if cl == Y:
            cell.number_format = MONEY
widths(ws, [9, 14, 10, 9, 10, 11, 12, 6, 7, 26, 38, 6, 10, 12, 11, 10, 22, 10, 10, 11, 9, 40,
            40, 12, 13, 11, 10, 14, 10, 14, 10, 12, 9])

# Mapping ----------------------------------------------------------------
ws = wb.create_sheet("Mapping")
header(ws, ["label_as_printed", "category", "vendor", "note"])
for i, row in enumerate(MAPPING, 2):
    for cidx, v in enumerate(row, 1):
        cell = ws.cell(row=i, column=cidx, value=v or None)
        cell.font = BLUE
        cell.fill = INPUT_FILL
widths(ws, [28, 12, 10, 50])

# Notes ------------------------------------------------------------------
ws = wb.create_sheet("Notes")
put(ws, "A1", "tracker_notes (paste Stage 3 output here; column E is a formula)", BOLD)
header(ws, ["project", "vendor", "category", "note", "key"], row=2)
for i, row in enumerate(TRACKER_NOTES, 3):
    for cidx, v in enumerate(row, 1):
        ws.cell(row=i, column=cidx, value=v).font = BLUE
for i in range(3, 3 + 50):
    put(ws, f"E{i}", f'=IF(A{i}="","",A{i}&"|"&B{i}&"|"&C{i})')
DN_ROW = 60
put(ws, f"A{DN_ROW}", "detail_notes (paste Stage 3 output here; column G is a formula)", BOLD)
header(ws, ["project", "vendor", "doc_id", "line_no", "standard_ref", "note", "key"], row=DN_ROW + 1)
ws.freeze_panes = "A3"
for i, row in enumerate(DETAIL_NOTES, DN_ROW + 2):
    for cidx, v in enumerate(row, 1):
        ws.cell(row=i, column=cidx, value=v).font = BLUE
for i in range(DN_ROW + 2, DN_ROW + 2 + 50):
    put(ws, f"G{i}", f'=IF(C{i}="","",C{i}&"|"&D{i})')
widths(ws, [14, 10, 12, 8, 12, 70, 30])
TN = f"Notes!$E$3:$E$52"; TN_NOTE = "Notes!$D$3:$D$52"
DNK = f"Notes!$G${DN_ROW+2}:$G${DN_ROW+51}"; DN_NOTE = f"Notes!$F${DN_ROW+2}:$F${DN_ROW+51}"

# Docs -------------------------------------------------------------------
ws = wb.create_sheet("Docs")
CATS = ["LABOUR_ALL", "INSTALL", "PRODUCT", "ALLOWANCE"]
dcols = (["seq", "row", "doc_id", "project", "vendor", "doc_type", "doc_date", "pages_total", "pages_read",
          "native_or_scan", "mapping_provided", "source_file"]
         + [f"{k}_{c}" for c in CATS for k in ("stated", "extracted", "check")]
         + ["stated_GRAND", "extracted_ALL", "check_GRAND", "dup_count", "review_reason", "review_seq"])
header(ws, dcols)
Dc = {n: col(n, dcols) for n in dcols}
Lr = lambda name: f"Lines!${L[name]}$2:${L[name]}${N_LINES}"
grp = lambda c: Lr("labour_group") if c == "LABOUR_ALL" else Lr("category")
TEXT_COLS = {"project", "vendor", "doc_type", "doc_id", "doc_date", "ref_quote_id", "label_as_printed",
             "description", "category", "status", "pages_read", "native_or_scan", "mapping_provided",
             "source_file", "review_reason"}
def idx(name, row):
    e = f"INDEX({Lr(name)},{row})"
    return e + '&""' if name in TEXT_COLS else e

for i in range(2, N_DOCS + 2):
    n = i - 1
    put(ws, f"{Dc['seq']}{i}", n)
    put(ws, f"{Dc['row']}{i}", f'=IFERROR(MATCH({Dc["seq"]}{i},{Lr("coverage_seq")},0),"")')
    row = f'{Dc["row"]}{i}'
    for name in ["doc_id", "project", "vendor", "doc_type", "doc_date", "pages_total", "pages_read",
                 "native_or_scan", "mapping_provided", "source_file"]:
        put(ws, f"{Dc[name]}{i}", f'=IF({row}="","",{idx(name, row)})')
    did = f'{Dc["doc_id"]}{i}'
    for c in CATS:
        st, ex, ck = (f'{Dc[f"{k}_{c}"]}{i}' for k in ("stated", "extracted", "check"))
        put(ws, st, f'=IF({did}="","",SUMIFS({Lr("amount")},{Lr("doc_id")},{did},{Lr("row_kind")},"total",{grp(c)},"{c}"))', fmt=MONEY)
        put(ws, ex, f'=IF({did}="","",SUMIFS({Lr("amount")},{Lr("doc_id")},{did},{Lr("row_kind")},"line",{grp(c)},"{c}"))', fmt=MONEY)
        put(ws, ck, f'=IF({did}="","",IF({st}=0,"no stated total",IF(ABS({st}-{ex})>{TOL},"MISMATCH","ok")))')
    sg, ea, cg = (f'{Dc[n_]}{i}' for n_ in ("stated_GRAND", "extracted_ALL", "check_GRAND"))
    put(ws, sg, f'=IF({did}="","",SUMIFS({Lr("amount")},{Lr("doc_id")},{did},{Lr("row_kind")},"total",{Lr("category")},"GRAND"))', fmt=MONEY)
    put(ws, ea, f'=IF({did}="","",SUMIFS({Lr("amount")},{Lr("doc_id")},{did},{Lr("row_kind")},"line"))', fmt=MONEY)
    put(ws, cg, f'=IF({did}="","",IF({sg}=0,"no stated total",IF(ABS({sg}-{ea})>{TOL},"MISMATCH","ok")))')
    dup = f'{Dc["dup_count"]}{i}'
    put(ws, dup, f'=IF({did}="","",COUNTIFS({Lr("doc_id")},{did},{Lr("row_kind")},"coverage"))')
    pt, pr, mp = (f'{Dc[n_]}{i}' for n_ in ("pages_total", "pages_read", "mapping_provided"))
    checks = "&".join(
        [f'IF({Dc[f"check_{c}"]}{i}="MISMATCH","lines do not sum to stated total ({c}); ","")' for c in CATS]
        + [f'IF({cg}="MISMATCH","lines do not sum to grand total; ","")',
           f'IF(OR({pr}="1-"&{pt},AND({pt}=1,{pr}="1")),"","pages not read; ")',
           f'IF({mp}="N","extracted without mapping table; ","")',
           f'IF({dup}>1,"document extracted twice; ","")'])
    rr = f'{Dc["review_reason"]}{i}'
    put(ws, rr, f'=IF({did}="","",IF(LEN({checks})>2,LEFT({checks},LEN({checks})-2),""))')
    put(ws, f'{Dc["review_seq"]}{i}', f'=IF({rr}<>"",MAX(${Dc["review_seq"]}$1:{Dc["review_seq"]}{i-1})+1,"")')
widths(ws, [5, 5, 10, 14, 10, 9, 11, 7, 7, 8, 8, 40] + [13, 13, 15] * 4 + [13, 13, 15, 6, 60, 8])

# Tracker ----------------------------------------------------------------
ws = wb.create_sheet("Tracker")
tcols = (["seq", "row", "project", "vendor"]
         + [f"{k}_{c}" for c in CATS for k in ("quoted", "invoiced", "variance", "variance_pct")]
         + ["invoice_count", "discrepancy", "notes"])
header(ws, tcols)
Tc = {n: col(n, tcols) for n in tcols}
Dr = lambda name: f"Docs!${Dc[name]}$2:${Dc[name]}${N_DOCS+1}"
for i in range(2, N_TRACK + 2):
    put(ws, f"{Tc['seq']}{i}", i - 1)
    put(ws, f"{Tc['row']}{i}", f'=IFERROR(MATCH({Tc["seq"]}{i},{Lr("pair_seq")},0),"")')
    row = f'{Tc["row"]}{i}'
    put(ws, f"{Tc['project']}{i}", f'=IF({row}="","",INDEX({Lr("project")},{row}))')
    put(ws, f"{Tc['vendor']}{i}", f'=IF({row}="","",INDEX({Lr("vendor")},{row}))')
    p, v = f'{Tc["project"]}{i}', f'{Tc["vendor"]}{i}'
    for c in CATS:
        q, inv, var, pct = (f'{Tc[f"{k}_{c}"]}{i}' for k in ("quoted", "invoiced", "variance", "variance_pct"))
        base = f'{Lr("amount")},{Lr("project")},{p},{Lr("vendor")},{v},{Lr("row_kind")},"line",{grp(c)},"{c}"'
        put(ws, q, f'=IF({p}="","",SUMIFS({base},{Lr("doc_type")},"quote"))', fmt=MONEY)
        put(ws, inv, f'=IF({p}="","",SUMIFS({base},{Lr("doc_type")},"invoice"))', fmt=MONEY)
        put(ws, var, f'=IF({p}="","",{inv}-{q})', fmt=MONEY)
        put(ws, pct, f'=IF({p}="","",IFERROR({var}/{q},""))', fmt=PCT)
    ic = f'{Tc["invoice_count"]}{i}'
    put(ws, ic, f'=IF({p}="","",COUNTIFS({Dr("project")},{p},{Dr("vendor")},{v},{Dr("doc_type")},"invoice"))')
    vars_ = ",".join(f'ABS({Tc[f"variance_{c}"]}{i})' for c in CATS)
    put(ws, f'{Tc["discrepancy"]}{i}', f'=IF({p}="","",IF({ic}=0,"N/A - pending",IF(MAX({vars_})>{TOL},"Y","N")))')
    notes = "&".join(
        f'IFERROR("{c}: "&INDEX({TN_NOTE},MATCH({p}&"|"&{v}&"|{c}",{TN},0))&" ","")' for c in CATS)
    put(ws, f'{Tc["notes"]}{i}', f'=IF({p}="","",TRIM({notes}))')
widths(ws, [5, 5, 14, 10] + [12, 12, 12, 9] * 4 + [8, 14, 90])

# Detail -----------------------------------------------------------------
ws = wb.create_sheet("Detail")
ecols = ["seq", "row", "project", "vendor", "doc_type", "doc_id", "ref_quote_id", "page", "line_no",
         "label_as_printed", "description", "category", "quote_amount", "invoice_amount", "variance",
         "status", "note"]
header(ws, ecols)
Ec = {n: col(n, ecols) for n in ecols}
for i in range(2, N_DETAIL + 2):
    put(ws, f"{Ec['seq']}{i}", i - 1)
    put(ws, f"{Ec['row']}{i}", f'=IFERROR(MATCH({Ec["seq"]}{i},{Lr("detail_seq")},0),"")')
    row = f'{Ec["row"]}{i}'
    for name in ["project", "vendor", "doc_type", "doc_id", "ref_quote_id", "page", "line_no",
                 "label_as_printed", "description", "category", "status"]:
        put(ws, f"{Ec[name]}{i}", f'=IF({row}="","",{idx(name, row)})')
    dt = f'{Ec["doc_type"]}{i}'
    qa, ia, va = (f'{Ec[n_]}{i}' for n_ in ("quote_amount", "invoice_amount", "variance"))
    put(ws, qa, f'=IF({row}="","",IF({dt}="quote",INDEX({Lr("amount")},{row}),""))', fmt=MONEY)
    put(ws, ia, f'=IF({row}="","",IF({dt}="quote",IF(INDEX({Lr("invoice_hits")},{row})=0,"",INDEX({Lr("invoice_amount")},{row})),INDEX({Lr("amount")},{row})))', fmt=MONEY)
    put(ws, va, f'=IF({row}="","",N({ia})-N({qa}))', fmt=MONEY)
    did, ln = f'{Ec["doc_id"]}{i}', f'{Ec["line_no"]}{i}'
    put(ws, f'{Ec["note"]}{i}', f'=IF({row}="","",IFERROR(INDEX({DN_NOTE},MATCH({did}&"|"&{ln},{DNK},0)),""))')
widths(ws, [5, 5, 14, 10, 9, 10, 12, 6, 7, 28, 40, 11, 13, 13, 12, 14, 60])

# Review -----------------------------------------------------------------
ws = wb.create_sheet("Review")
rcols = ["seq", "doc_id", "page", "line_no", "label_as_printed", "reason"]
header(ws, rcols)
Rc = {n: col(n, rcols) for n in rcols}
line_count = f'COUNT({Lr("review_seq")})'
for i in range(2, N_REVIEW + 2):
    s = f'{Rc["seq"]}{i}'
    put(ws, s, i - 1)
    lrow = f'MATCH({s},{Lr("review_seq")},0)'
    drow = f'MATCH({s}-{line_count},{Dr("review_seq")},0)'
    def pick(line_expr, doc_expr):
        return f'=IF({s}<={line_count},IFERROR({line_expr},""),IFERROR({doc_expr},""))'
    put(ws, f'{Rc["doc_id"]}{i}', pick(f'INDEX({Lr("doc_id")},{lrow})&""', f'INDEX({Dr("doc_id")},{drow})&""'))
    put(ws, f'{Rc["page"]}{i}', pick(f'INDEX({Lr("page")},{lrow})', '""'))
    put(ws, f'{Rc["line_no"]}{i}', pick(f'INDEX({Lr("line_no")},{lrow})', '""'))
    put(ws, f'{Rc["label_as_printed"]}{i}', pick(f'INDEX({Lr("label_as_printed")},{lrow})&""', '""'))
    put(ws, f'{Rc["reason"]}{i}', pick(f'INDEX({Lr("review_reason")},{lrow})&""', f'INDEX({Dr("review_reason")},{drow})&""'))
widths(ws, [5, 10, 6, 7, 30, 70])

wb.calculation.fullCalcOnLoad = True  # no cached values are written; Excel recalculates on open
wb.save(OUT)
print("wrote", OUT.name, "data rows in Lines:", last_data_row - 1)

# ---------------------------------------------------------------- sample CSVs (what Stage 1 returns)
for doc_id, rows in DOCS:
    first = rows[0]
    name = f"{first[1]}__{first[2]}__{first[3]}__{doc_id}.csv"
    with open(CSV_DIR / name, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(CSV_COLS)
        for row in rows:
            w.writerow(row)
    print("wrote csv/", name)
